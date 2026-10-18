from __future__ import annotations
import os
import fitz
import docx

FIXTURES_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "fixtures"))
os.makedirs(FIXTURES_DIR, exist_ok=True)


def create_clean_1page_pdf() -> str:
    path = os.path.join(FIXTURES_DIR, "clean_1page_resume.pdf")
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)  # Standard A4

    text = """Alex Rivera
alex.rivera@example.com | (555) 123-4567 | San Francisco, CA
https://github.com/alexrivera | https://linkedin.com/in/alexrivera

EDUCATION
State University of California
Bachelor of Science in Computer Science | 2020 - 2024
GPA: 3.8 / 4.0

TECHNICAL SKILLS
Languages: Python, TypeScript, JavaScript, SQL, C++, HTML, CSS
Frameworks & Libraries: React, FastAPI, Node.js, Express, Next.js
Tools & Platforms: Docker, AWS, Git, GitHub, PostgreSQL, Redis, Linux

WORK EXPERIENCE
Software Engineering Intern | Acme Cloud Solutions | Summer 2023
- Developed REST API endpoints using FastAPI and PostgreSQL, serving 50k daily active users.
- Containerized microservices using Docker and automated CI/CD deployment pipelines.
- Improved query response times by 35% through indexing and caching with Redis.

KEY PROJECTS
Readiness Copilot | https://github.com/alexrivera/readiness-copilot | https://readiness.dev
- Built full-stack placement readiness platform with Python, FastAPI, and React.
- Designed structured JSON review workflows and automated resume benchmark checks.

Distributed Task Queue | https://github.com/alexrivera/task-queue
- Implemented high-throughput asynchronous task worker in Python using Redis and Docker.

ACHIEVEMENTS
- 1st Place Winner at HackCal 2023 out of 120 teams.
- AWS Certified Cloud Practitioner (2023).
"""
    # Insert text into the page
    rect = fitz.Rect(40, 40, 555, 800)
    page.insert_textbox(rect, text, fontsize=9.5, fontname="helv")
    doc.save(path)
    doc.close()
    return path


def create_messy_multipage_pdf() -> str:
    path = os.path.join(FIXTURES_DIR, "messy_multipage_resume.pdf")
    doc = fitz.open()

    # Page 1
    p1 = doc.new_page(width=595, height=842)
    p1_text = """JORDAN SMITH
jordan.smith@university.edu | (555) 987-6543
https://github.com/jordansmith | https://linkedin.com/in/jordansmith-engineer

PROFESSIONAL SUMMARY
Senior computer science undergraduate with 3 internships and deep interest in distributed systems,
machine learning, and backend scalability. Proficient in Python, Go, and Kubernetes.

EDUCATION
Metropolitan Institute of Technology
Master of Science in Software Engineering, Expected 2025
Bachelor of Science in Computer Science, 2023

SKILLS
Python, Go, Java, Kubernetes, Docker, AWS, GCP, PostgreSQL, MongoDB, Redis, Kafka, CI/CD, Linux
"""
    p1.insert_textbox(fitz.Rect(40, 40, 555, 800), p1_text, fontsize=10, fontname="helv")

    # Page 2
    p2 = doc.new_page(width=595, height=842)
    p2_text = """JORDAN SMITH - PAGE 2

WORK EXPERIENCE
Distributed Systems Intern | Global Data Corp | Jan 2023 - Aug 2023
- Designed streaming telemetry ingestion pipeline processing 10M records/day with Kafka and Go.
- Orchestrated 45-node Kubernetes cluster deployments across AWS and GCP.
- Implemented automated monitoring and alerting with Prometheus and Grafana.

Backend Development Intern | TechFlow Dynamics | May 2022 - Dec 2022
- Maintained core microservices using Python and PostgreSQL.
- Optimized database query bottlenecks reducing p99 latency from 450ms to 85ms.
"""
    p2.insert_textbox(fitz.Rect(40, 40, 555, 800), p2_text, fontsize=10, fontname="helv")

    # Page 3
    p3 = doc.new_page(width=595, height=842)
    p3_text = """JORDAN SMITH - PAGE 3

PROJECTS
Event Stream Processing Engine | https://github.com/jordansmith/stream-engine
- High performance pub-sub engine written in Go with Raft consensus protocol.
- Benchmarked at 150k messages per second on single-node benchmarks.

Cloud Resource Allocator | https://github.com/jordansmith/allocator | https://allocator-demo.org
- Intelligent VM scheduling simulator leveraging Python, Pandas, and Scikit-Learn.

AWARDS & CERTIFICATIONS
- Published paper on lightweight consensus in IEEE Student Symposium 2023.
- Certified Kubernetes Administrator (CKA).
"""
    p3.insert_textbox(fitz.Rect(40, 40, 555, 800), p3_text, fontsize=10, fontname="helv")

    doc.save(path)
    doc.close()
    return path


def create_sample_docx() -> str:
    path = os.path.join(FIXTURES_DIR, "sample_resume.docx")
    doc = docx.Document()

    # Title / Header
    h = doc.add_heading("Taylor Morgan", level=0)
    p_links = doc.add_paragraph("Email: taylor.morgan@example.com | Phone: 555-444-3333\n")
    p_links.add_run("GitHub: https://github.com/taylormorgan | LinkedIn: https://linkedin.com/in/taylormorgan")

    # Education
    doc.add_heading("Education", level=1)
    p_edu = doc.add_paragraph()
    p_edu.add_run("Tech University\n").bold = True
    p_edu.add_run("Bachelor of Science in Computer Science, Graduated 2024\n")
    p_edu.add_run("Relevant Coursework: Data Structures, Operating Systems, Web Systems")

    # Skills
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph(
        "Languages: Java, Python, TypeScript, SQL\n"
        "Frameworks: Spring Boot, React, Node.js, Express\n"
        "Databases & Tools: MySQL, PostgreSQL, Docker, Git, REST API"
    )

    # Experience
    doc.add_heading("Experience", level=1)
    p_exp = doc.add_paragraph()
    p_exp.add_run("Junior Full Stack Developer | Horizon Labs (2023 - 2024)\n").bold = True
    p_exp.add_run("- Built responsive dashboards using React and TypeScript.\n")
    p_exp.add_run("- Created backend REST APIs with Java and Spring Boot connecting to MySQL.\n")
    p_exp.add_run("- Wrote unit and integration tests achieving 85% test coverage.")

    # Projects
    doc.add_heading("Projects", level=1)
    p_proj = doc.add_paragraph()
    p_proj.add_run("E-Commerce Portal | https://github.com/taylormorgan/ecommerce-portal | https://taylor-shop.com\n").bold = True
    p_proj.add_run("- Full-stack e-commerce store with user authentication, cart, and payment processing.\n")
    p_proj.add_run("- Tech stack: React, Node.js, MySQL, Docker.")

    # Achievements
    doc.add_heading("Achievements", level=1)
    doc.add_paragraph(
        "- Hackathon Finalist at TechHacks 2023.\n"
        "- President of the University ACM Student Chapter."
    )

    doc.save(path)
    return path


def create_corrupt_file() -> str:
    path = os.path.join(FIXTURES_DIR, "corrupt_resume.pdf")
    with open(path, "wb") as f:
        # Write corrupted header and invalid truncated garbage bytes
        f.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\nThis is corrupt junk that is not a valid PDF structure at all!\x00\xff\xfe\x12\x34\x56\x78")
    return path


def create_resume_good() -> str:
    path = os.path.join(FIXTURES_DIR, "resume_good.pdf")
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)  # Standard A4

    text = """Alex Rivera
alex.rivera@example.com | (555) 123-4567 | San Francisco, CA
https://github.com/alexrivera | https://leetcode.com/u/alexrivera | https://linkedin.com/in/alexrivera

EDUCATION
State University of California
Bachelor of Science in Computer Science | Sep 2020 - May 2024
GPA: 3.8 / 4.0

TECHNICAL SKILLS
Languages: Python, TypeScript, JavaScript, SQL, C++, HTML, CSS
Frameworks & Libraries: React, FastAPI, Node.js, Express, Next.js
Tools & Platforms: Docker, AWS, Git, GitHub, PostgreSQL, Redis, Linux

WORK EXPERIENCE
Software Engineering Intern | Acme Cloud Solutions | Jun 2023 - Aug 2023
- Developed REST API endpoints using FastAPI and PostgreSQL, serving 50k daily active users.
- Containerized microservices using Docker and automated CI/CD deployment pipelines.
- Improved query response times by 35% through indexing and caching with Redis.

KEY PROJECTS
Readiness Copilot | https://github.com/alexrivera/readiness-copilot | https://readiness.dev
- Built full-stack placement readiness platform with Python, FastAPI, and React, serving 500 active students.
- Tech Stack: Python, FastAPI, React, PostgreSQL, Docker
- Designed structured JSON review workflows and automated resume benchmark checks.

Distributed Task Queue | https://github.com/alexrivera/task-queue | https://taskqueue.io
- Implemented high-throughput asynchronous task worker in Python, processing 150k messages per second.
- Tech Stack: Python, Redis, Docker, RabbitMQ

ACHIEVEMENTS
- 1st Place Winner at HackCal 2023 out of 120 teams.
- AWS Certified Cloud Practitioner (Sep 2023).
- President of University Open Source & Competitive Programming Club.
"""
    rect = fitz.Rect(40, 40, 555, 800)
    page.insert_textbox(rect, text, fontsize=9.5, fontname="helv")
    doc.save(path)
    doc.close()
    return path


def create_resume_2page() -> str:
    path = os.path.join(FIXTURES_DIR, "resume_2page.pdf")
    doc = fitz.open()

    # Page 1
    p1 = doc.new_page(width=595, height=842)
    p1_text = """Alex Rivera
alex.rivera@example.com | (555) 123-4567 | San Francisco, CA
https://github.com/alexrivera | https://leetcode.com/u/alexrivera

EDUCATION
State University of California
Bachelor of Science in Computer Science | Sep 2020 - May 2024

TECHNICAL SKILLS
Languages: Python, TypeScript, JavaScript, SQL
Frameworks: React, FastAPI, Node.js, Express
Tools: Docker, AWS, Git, GitHub, PostgreSQL, Redis

WORK EXPERIENCE
Software Engineering Intern | Acme Cloud Solutions | Jun 2023 - Aug 2023
- Developed REST API endpoints using FastAPI and PostgreSQL, serving 50k daily active users.
- Containerized microservices using Docker and automated CI/CD deployment pipelines.
- Improved query response times by 35% through indexing and caching with Redis.
"""
    p1.insert_textbox(fitz.Rect(40, 40, 555, 800), p1_text, fontsize=10, fontname="helv")

    # Page 2
    p2 = doc.new_page(width=595, height=842)
    p2_text = """Alex Rivera - Page 2

PROJECTS
Readiness Copilot | https://github.com/alexrivera/readiness-copilot | https://readiness.dev
- Built full-stack placement readiness platform with Python, FastAPI, and React, serving 500 students.
- Tech Stack: Python, FastAPI, React, PostgreSQL

Distributed Task Queue | https://github.com/alexrivera/task-queue | https://taskqueue.io
- Implemented high-throughput asynchronous task worker in Python handling 150k messages per second.
- Tech Stack: Python, Redis, Docker

ACHIEVEMENTS
- 1st Place Winner at HackCal 2023 out of 120 teams.
- AWS Certified Cloud Practitioner.
"""
    p2.insert_textbox(fitz.Rect(40, 40, 555, 800), p2_text, fontsize=10, fontname="helv")

    doc.save(path)
    doc.close()
    return path


def create_resume_missing_links_typos() -> str:
    path = os.path.join(FIXTURES_DIR, "resume_missing_links_typos.pdf")
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)

    # Note: Missing LeetCode link, planted typos: 'managment', 'expereince', 'softare', weak bullets without action verbs or numbers
    text = """Alex Rivera
alex.rivera@example.com | (555) 123-4567 | San Francisco, CA
https://linkedin.com/in/alexrivera

EDUCATION
State University of California
Bachelor of Science in Computer Science | 2020 - 2024

TECHNICAL SKILLS
Languages: Python, TypeScript, JavaScript, SQL
Frameworks: React, FastAPI, Node.js

WORK EXPERIENCE
Software Developer Intern | Acme Cloud Solutions | Summer 2023
- Responsable for managment of softare expereince and databases.
- Worked on stuff for the website and assisted with team tasks.
- Helped the team with various daily engineering issues.

PROJECTS
Readiness Platform | https://github.com/alexrivera/readiness-platform
- A web application for students.

ACHIEVEMENTS
- Attended HackCal hackathon.
"""
    rect = fitz.Rect(40, 40, 555, 800)
    page.insert_textbox(rect, text, fontsize=10, fontname="helv")
    doc.save(path)
    doc.close()
    return path


def create_resume_image_only() -> str:
    path = os.path.join(FIXTURES_DIR, "resume_image_only.pdf")

    # Step 1: Create a temporary document with resume text
    temp_doc = fitz.open()
    temp_page = temp_doc.new_page(width=595, height=842)
    temp_page.insert_textbox(
        fitz.Rect(40, 40, 555, 800),
        "Alex Rivera\nScanned Resume Document\nEducation: CS Degree\nExperience: Intern at Tech Corp",
        fontsize=12,
        fontname="helv",
    )
    # Render the page to a raster pixmap
    pix = temp_page.get_pixmap(dpi=150)
    img_bytes = pix.tobytes("png")
    temp_doc.close()

    # Step 2: Create image-only PDF inserting the raster image with ZERO text stream
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(fitz.Rect(0, 0, 595, 842), stream=img_bytes)
    doc.save(path)
    doc.close()
    return path


def create_linkedin_fixtures():
    # Helper for avatar png
    temp_doc = fitz.open()
    temp_p = temp_doc.new_page(width=100, height=100)
    temp_p.draw_circle(fitz.Point(50, 50), 45, color=(0.1, 0.4, 0.8), fill=(0.2, 0.5, 0.9))
    pix = temp_p.get_pixmap()
    avatar_png = pix.tobytes("png")
    temp_doc.close()

    # 1. Strong LinkedIn Profile
    strong_txt = """Alex Rivera
Full Stack Software Engineer | React, TypeScript, Python & FastAPI | Building Scalable Systems
[Profile Photo: Yes]

About
I am a results-driven Full Stack Software Engineer experienced in designing and delivering high-performance web applications and distributed backends. Proficient in Python, FastAPI, React, TypeScript, PostgreSQL, and Docker, with a track record of building production-grade services. Passionate about system design, clean API architecture, and collaborating on high-impact products.

Experience
Software Engineering Intern at Acme Cloud Solutions
Summer 2023 - Present
Developed REST API endpoints using FastAPI and PostgreSQL, serving 50k daily active users. Containerized microservices using Docker.

Projects
Readiness Copilot
Built full-stack placement readiness platform with Python, FastAPI, and React.

Skills
Python, TypeScript, JavaScript, React, FastAPI, Node.js, PostgreSQL, Docker, Git, Redis, SQL
"""
    strong_txt_path = os.path.join(FIXTURES_DIR, "linkedin_strong.txt")
    with open(strong_txt_path, "w", encoding="utf-8") as f:
        f.write(strong_txt)

    strong_pdf_path = os.path.join(FIXTURES_DIR, "linkedin_strong.pdf")
    doc_strong = fitz.open()
    p_strong = doc_strong.new_page(width=595, height=842)
    p_strong.insert_image(fitz.Rect(40, 40, 110, 110), stream=avatar_png)
    p_strong.insert_textbox(fitz.Rect(40, 120, 555, 800), strong_txt, fontsize=10, fontname="helv")
    doc_strong.save(strong_pdf_path)
    doc_strong.close()

    # 2. Weak LinkedIn Profile
    weak_txt = """Jordan Doe
Student at MSOT
[Profile Photo: Missing]

About
I like coding and web stuff.

Experience
Coding Club Member at University

Skills
HTML, CSS
"""
    weak_txt_path = os.path.join(FIXTURES_DIR, "linkedin_weak.txt")
    with open(weak_txt_path, "w", encoding="utf-8") as f:
        f.write(weak_txt)

    weak_pdf_path = os.path.join(FIXTURES_DIR, "linkedin_weak.pdf")
    doc_weak = fitz.open()
    p_weak = doc_weak.new_page(width=595, height=842)
    # No avatar image inserted
    p_weak.insert_textbox(fitz.Rect(40, 50, 555, 800), weak_txt, fontsize=10, fontname="helv")
    doc_weak.save(weak_pdf_path)
    doc_weak.close()

    # 3. Contradictory LinkedIn Profile (claims VP of Engineering vs student intern)
    contra_txt = """Alex Rivera
Vice President of Engineering | Global Tech Strategy | Enterprise Architect
[Profile Photo: Yes]

About
Executive technology leader directing global infrastructure, enterprise architecture, and distributed engineering initiatives. Over a decade of leading world-class teams and scaling high-availability systems.

Experience
VP of Engineering at Acme Cloud Solutions
2021 - Present
Directed architectural strategy, budget allocation, and engineering staffing across global distributed cloud systems.

Skills
Python, FastAPI, React, TypeScript, PostgreSQL, Docker, Git, Cloud Architecture
"""
    contra_txt_path = os.path.join(FIXTURES_DIR, "linkedin_contradictory.txt")
    with open(contra_txt_path, "w", encoding="utf-8") as f:
        f.write(contra_txt)

    contra_pdf_path = os.path.join(FIXTURES_DIR, "linkedin_contradictory.pdf")
    doc_contra = fitz.open()
    p_contra = doc_contra.new_page(width=595, height=842)
    p_contra.insert_image(fitz.Rect(40, 40, 110, 110), stream=avatar_png)
    p_contra.insert_textbox(fitz.Rect(40, 120, 555, 800), contra_txt, fontsize=10, fontname="helv")
    doc_contra.save(contra_pdf_path)
    doc_contra.close()

    return {
        "strong_txt": strong_txt_path,
        "strong_pdf": strong_pdf_path,
        "weak_txt": weak_txt_path,
        "weak_pdf": weak_pdf_path,
        "contra_txt": contra_txt_path,
        "contra_pdf": contra_pdf_path,
    }


def generate_all_fixtures():
    p1 = create_clean_1page_pdf()
    p2 = create_messy_multipage_pdf()
    p3 = create_sample_docx()
    p4 = create_corrupt_file()
    r_good = create_resume_good()
    r_2page = create_resume_2page()
    r_flawed = create_resume_missing_links_typos()
    r_img = create_resume_image_only()
    li_fixtures = create_linkedin_fixtures()
    return {
        "clean_pdf": p1,
        "multipage_pdf": p2,
        "docx": p3,
        "corrupt": p4,
        "resume_good": r_good,
        "resume_2page": r_2page,
        "resume_missing_links_typos": r_flawed,
        "resume_image_only": r_img,
        **li_fixtures,
    }



if __name__ == "__main__":
    fixtures = generate_all_fixtures()
    print("Generated fixtures successfully:")
    for k, v in fixtures.items():
        print(f"  - {k}: {v} (size: {os.path.getsize(v)} bytes)")

