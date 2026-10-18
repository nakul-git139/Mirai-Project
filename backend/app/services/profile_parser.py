from __future__ import annotations
import io
import re
from typing import Dict, Any, List, Optional
import fitz  # PyMuPDF
import docx

# Curated skills taxonomy for deterministic, benchmark-level extraction
COMMON_SKILLS = [
    # Languages
    "Python", "JavaScript", "TypeScript", "Java", "C++", "C", "C#", "Go", "Rust",
    "Ruby", "PHP", "Swift", "Kotlin", "SQL", "R", "Bash", "Shell", "HTML", "CSS",
    # Frameworks & Libraries
    "React", "Next.js", "Node.js", "Express", "Vue", "Angular", "Django", "FastAPI",
    "Flask", "Spring", "Spring Boot", "PyTorch", "TensorFlow", "Scikit-Learn",
    "Pandas", "NumPy", "Tailwind CSS", "Bootstrap", "GraphQL", "Redux",
    # Cloud, DevOps & Tools
    "Docker", "Kubernetes", "AWS", "GCP", "Azure", "Git", "GitHub", "Linux",
    "CI/CD", "PostgreSQL", "MySQL", "SQLite", "MongoDB", "Redis", "Kafka",
    "Elasticsearch", "REST API", "Microservices"
]

SECTION_KEYWORDS = {
    "education": ["education", "academic background", "academics", "qualifications"],
    "experience": ["experience", "work experience", "employment", "professional experience", "internships", "work history"],
    "projects": ["projects", "personal projects", "academic projects", "key projects"],
    "skills": ["skills", "technical skills", "technologies", "proficiencies", "core competencies", "tools"],
    "achievements": ["achievements", "certifications", "awards", "honors", "hackathons", "extracurricular"],
    "summary": ["summary", "profile", "about me", "professional summary", "objective"]
}


def extract_text_and_links_from_pdf(file_bytes: bytes) -> Dict[str, Any]:
    """
    Extract text, page count, and embedded hyperlinks from a PDF using PyMuPDF (fitz).
    Raises ValueError on corrupted or invalid PDFs.
    """
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise ValueError(f"Invalid or corrupt PDF file: {str(e)}")

    if len(doc) == 0:
        raise ValueError("The provided PDF file has 0 pages or contains no readable content.")

    page_count = len(doc)
    pages_text = []
    links = set()

    for page_num in range(page_count):
        page = doc[page_num]
        text = page.get_text("text")
        pages_text.append(text)

        # Extract embedded hyperlink annotations
        for link in page.get_links():
            uri = link.get("uri")
            if uri and uri.startswith(("http://", "https://", "mailto:")):
                links.add(uri)

    doc.close()
    full_text = "\n\n".join(pages_text)

    # Also detect plain-text URLs with regex
    url_pattern = re.compile(r"https?://[a-zA-Z0-9_\-./?%&=+#@:~]+")
    for match in url_pattern.finditer(full_text):
        links.add(match.group(0).rstrip(".,;)"))

    return {
        "text": full_text,
        "page_count": page_count,
        "links": sorted(list(links)),
    }


def extract_text_and_links_from_docx(file_bytes: bytes) -> Dict[str, Any]:
    """
    Extract text, estimated page count, and hyperlinks from a DOCX file using python-docx.
    Raises ValueError on corrupted or invalid DOCX files.
    """
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Invalid or corrupt DOCX file: {str(e)}")

    paragraphs_text = []
    links = set()

    # Extract paragraph text and hyperlinks from document XML relationships
    for p in doc.paragraphs:
        paragraphs_text.append(p.text)
        # Search XML for hyperlinks
        for rId in p._p.xpath(".//w:hyperlink/@r:id"):
            try:
                rel = doc.part.rels[rId]
                target = rel.target_ref
                if target.startswith(("http://", "https://", "mailto:")):
                    links.add(target)
            except Exception:
                pass

    # Extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                paragraphs_text.append(" | ".join(row_text))

    full_text = "\n".join(paragraphs_text).strip()
    if not full_text:
        raise ValueError("The provided DOCX document contains no readable text.")

    # Plain text URLs
    url_pattern = re.compile(r"https?://[a-zA-Z0-9_\-./?%&=+#@:~]+")
    for match in url_pattern.finditer(full_text):
        links.add(match.group(0).rstrip(".,;)"))

    # Estimate page count: typical resume pages have ~450 words
    word_count = len(full_text.split())
    estimated_pages = max(1, (word_count + 400) // 450)

    return {
        "text": full_text,
        "page_count": estimated_pages,
        "links": sorted(list(links)),
    }


def identify_sections(text: str) -> Dict[str, str]:
    """
    Segments resume text into known canonical sections based on standard headings.
    """
    lines = [line.strip() for line in text.splitlines()]
    sections: Dict[str, List[str]] = {}
    current_section = "header"
    sections[current_section] = []

    for line in lines:
        if not line:
            continue

        clean_line = line.strip().lower().rstrip(":")
        matched_section = None

        # Check if line looks like a header (short, matches section keywords)
        if len(clean_line) <= 40:
            for sec_name, keywords in SECTION_KEYWORDS.items():
                if any(kw == clean_line or clean_line.startswith(kw + " ") or clean_line.endswith(" " + kw) for kw in keywords):
                    matched_section = sec_name
                    break

        if matched_section:
            current_section = matched_section
            if current_section not in sections:
                sections[current_section] = []
        else:
            sections[current_section].append(line)

    return {k: "\n".join(v).strip() for k, v in sections.items() if v}


def extract_skills_from_text(text: str) -> List[str]:
    """
    Extracts recognized skills matching our standard taxonomy.
    """
    found_skills = set()
    text_lower = f" {text.lower()} "

    for skill in COMMON_SKILLS:
        # Regex matching word boundary or symbols like C++ / C# / .NET
        escaped = re.escape(skill.lower())
        pattern = rf"(?<![\w#+]){escaped}(?![\w#+])"
        if re.search(pattern, text_lower):
            found_skills.add(skill)

    return sorted(list(found_skills))


def extract_projects_from_text(projects_text: str, all_links: List[str]) -> List[Dict[str, Any]]:
    """
    Parses project blocks from project section text, associating repo and live links.
    """
    if not projects_text:
        return []

    blocks = re.split(r"\n(?=[A-Z0-9][A-Za-z0-9\s\-]{2,40}(?:\s*\||\s*\(|\s*[-–]|\n|$))", projects_text)
    projects = []

    for idx, block in enumerate(blocks):
        block_clean = block.strip()
        if not block_clean:
            continue

        lines = block_clean.splitlines()
        first_line = lines[0].strip()
        title = re.split(r"\||\(|\-|\–", first_line)[0].strip()
        if not title or len(title) > 60:
            title = f"Project {idx + 1}"

        # Find repo or live links inside this block
        repo_url = None
        live_url = None

        block_urls = re.findall(r"https?://[a-zA-Z0-9_\-./?%&=+#@:~]+", block_clean)
        for url in block_urls:
            url_clean = url.rstrip(".,;)")
            if "github.com" in url_clean.lower() and not repo_url:
                repo_url = url_clean
            elif not live_url and not "github.com" in url_clean.lower():
                live_url = url_clean

        # If repo link wasn't inline, check overall links
        if not repo_url:
            for l in all_links:
                if "github.com" in l.lower() and title.lower().replace(" ", "") in l.lower():
                    repo_url = l
                    break

        # Extract skills in this project block
        block_skills = extract_skills_from_text(block_clean)
        description = "\n".join(lines[1:]).strip() if len(lines) > 1 else block_clean

        projects.append({
            "id": f"proj-{idx + 1}",
            "title": title,
            "repo_url": repo_url,
            "live_url": live_url,
            "description": description[:500],
            "tech_stack": block_skills,
        })

    return projects


def extract_education_from_text(education_text: str) -> List[Dict[str, Any]]:
    """
    Parses education blocks into degree, institution, and year.
    """
    if not education_text:
        return []

    lines = [l.strip() for l in education_text.splitlines() if l.strip()]
    education_list = []

    curr_entry: Dict[str, Any] = {"institution": "", "degree": "", "year": "", "details": ""}
    year_pattern = re.compile(r"\b(19\d{2}|20\d{2})\b")

    for line in lines:
        years = year_pattern.findall(line)
        if any(term in line.lower() for term in ["university", "college", "institute", "school"]):
            if curr_entry["institution"]:
                education_list.append(curr_entry)
                curr_entry = {"institution": "", "degree": "", "year": "", "details": ""}
            curr_entry["institution"] = line
            if years:
                curr_entry["year"] = years[-1]
        elif any(deg in line.lower() for deg in ["bachelor", "master", "phd", "b.s", "m.s", "b.tech", "m.tech", "degree"]):
            curr_entry["degree"] = line
            if years and not curr_entry["year"]:
                curr_entry["year"] = years[-1]
        else:
            if curr_entry["details"]:
                curr_entry["details"] += " | " + line
            else:
                curr_entry["details"] = line
            if years and not curr_entry["year"]:
                curr_entry["year"] = years[-1]

    if curr_entry["institution"] or curr_entry["degree"]:
        education_list.append(curr_entry)

    return education_list


def parse_resume_file(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Full pipeline to parse PDF or DOCX resume into structured data.
    """
    filename_lower = filename.lower()
    if filename_lower.endswith(".pdf"):
        extracted = extract_text_and_links_from_pdf(file_bytes)
    elif filename_lower.endswith((".docx", ".doc")):
        extracted = extract_text_and_links_from_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file format for '{filename}'. Please upload a PDF or DOCX file.")

    text = extracted["text"]
    page_count = extracted["page_count"]
    links = extracted["links"]

    sections = identify_sections(text)

    # Extract skills prioritizing skills section, then full text
    skills_text = sections.get("skills", "") + "\n" + text
    skills = extract_skills_from_text(skills_text)

    # Extract projects
    projects_text = sections.get("projects", "")
    projects = extract_projects_from_text(projects_text, links)

    # Extract education
    education_text = sections.get("education", "")
    education = extract_education_from_text(education_text)

    return {
        "filename": filename,
        "page_count": page_count,
        "sections": {k: v for k, v in sections.items() if k != "header"},
        "skills": skills,
        "projects": projects,
        "education": education,
        "links": links,
        "raw_text": text,
    }


def parse_linkedin_text(text: str) -> Dict[str, Any]:
    """
    Extracts LinkedIn profile fields (headline, about, experience, skills) from pasted text.
    """
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if not lines:
        raise ValueError("LinkedIn text is empty. Please paste your LinkedIn profile text.")

    headline = ""
    about = ""
    skills = extract_skills_from_text(text)
    experience_list = []

    # Partition sections by matching dedicated header lines
    current_sec = "top"
    sec_content: Dict[str, List[str]] = {"top": [], "about": [], "experience": [], "skills": []}

    for line in lines:
        clean_header = line.strip().lower().rstrip(":")
        if len(clean_header) <= 35:
            if clean_header in ["about", "summary", "about me", "professional summary"]:
                current_sec = "about"
                continue
            elif clean_header in ["experience", "work experience", "work history", "positions", "employment"]:
                current_sec = "experience"
                continue
            elif clean_header in ["skills", "top skills", "endorsements", "skills & endorsements", "proficiencies"]:
                current_sec = "skills"
                continue

        sec_content[current_sec].append(line)

    # Headline is typically the 2nd line of top (after name) or top line if descriptive
    top_lines = sec_content.get("top", [])
    if len(top_lines) > 1:
        headline = top_lines[1]
    elif len(top_lines) == 1:
        headline = top_lines[0]

    about = "\n".join(sec_content.get("about", [])).strip()

    # Parse simple experience entries
    exp_lines = sec_content.get("experience", [])
    if exp_lines:
        curr_exp: Dict[str, Any] = {"title": "", "company": "", "period": "", "description": ""}
        for line in exp_lines:
            if " at " in line.lower() or " | " in line:
                if curr_exp["title"]:
                    experience_list.append(curr_exp)
                    curr_exp = {"title": "", "company": "", "period": "", "description": ""}
                parts = re.split(r"\bat\b|\|", line, flags=re.IGNORECASE)
                curr_exp["title"] = parts[0].strip()
                if len(parts) > 1:
                    curr_exp["company"] = parts[1].strip()
            else:
                if not curr_exp["title"]:
                    curr_exp["title"] = line
                elif not curr_exp["company"]:
                    curr_exp["company"] = line
                else:
                    curr_exp["description"] += (" " + line).strip()
        if curr_exp["title"]:
            experience_list.append(curr_exp)

    return {
        "headline": headline[:250],
        "about": about,
        "experience": experience_list,
        "skills": skills,
        "source": "pasted",
    }


def parse_linkedin_pdf(file_bytes: bytes) -> Dict[str, Any]:
    """
    Parses standard LinkedIn PDF export into structured fields.
    """
    extracted = extract_text_and_links_from_pdf(file_bytes)
    text = extracted["text"]
    parsed = parse_linkedin_text(text)
    parsed["source"] = "pdf"
    return parsed


def extract_profile_data_from_chat_text(message: str) -> Dict[str, Any]:
    """
    Heuristically extracts GitHub username/URL, LeetCode username/URL,
    and project links from informal chat messages.
    """
    result: Dict[str, Any] = {}

    # GitHub patterns (matches with or without https://)
    github_url_match = re.search(r"(?:https?://)?(?:www\.)?github\.com/([a-zA-Z0-9_-]+)(?:/[a-zA-Z0-9_\-.]+)?", message, re.IGNORECASE)
    github_user_match = re.search(r"(?:my\s+)?github(?:\s+(?:is|username|profile|handle))?[:\s]+@?([a-zA-Z0-9_-]+)(?!\.com)", message, re.IGNORECASE)

    if github_url_match:
        username = github_url_match.group(1)
        result["github_username"] = username
        result["github_url"] = f"https://github.com/{username}"
    elif github_user_match:
        username = github_user_match.group(1).lstrip("@")
        if username.lower() not in ["http", "https", "www", "url"]:
            result["github_username"] = username
            result["github_url"] = f"https://github.com/{username}"

    # LeetCode patterns (matches with or without https://)
    leetcode_url_match = re.search(r"(?:https?://)?(?:www\.)?leetcode\.com/(?:u/)?([a-zA-Z0-9_-]+)/?", message, re.IGNORECASE)
    leetcode_user_match = re.search(r"(?:my\s+)?leetcode(?:\s+(?:is|username|profile|handle))?[:\s]+@?([a-zA-Z0-9_-]+)(?!\.com)", message, re.IGNORECASE)

    if leetcode_url_match:
        username = leetcode_url_match.group(1)
        result["leetcode_username"] = username
        result["leetcode_url"] = f"https://leetcode.com/u/{username}/"
    elif leetcode_user_match:
        username = leetcode_user_match.group(1).lstrip("@")
        if username.lower() not in ["http", "https", "www", "url"]:
            result["leetcode_username"] = username
            result["leetcode_url"] = f"https://leetcode.com/u/{username}/"

    # Project mentions (e.g., project called X at github.com/...)
    project_match = re.search(r"(?:project|built|created)\s+([A-Za-z0-9_\-\s]{3,30})\s*(?:at|repo|live)?\s*(https?://[^\s]+)?", message, re.IGNORECASE)
    if project_match:
        title = project_match.group(1).strip()
        url = project_match.group(2) if project_match.group(2) else None
        result["project"] = {
            "title": title,
            "repo_url": url if url and "github.com" in url.lower() else None,
            "live_url": url if url and "github.com" not in url.lower() else None,
            "description": f"Shared via chat: {message.strip()}",
            "tech_stack": extract_skills_from_text(message),
        }

    return result
